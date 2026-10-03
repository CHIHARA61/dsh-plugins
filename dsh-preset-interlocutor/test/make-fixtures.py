#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""生成 OOXML 抽取测试用的 fixtures（可复现）。

用法：
    <bundled-python> test/make-fixtures.py [输出目录]

默认输出到 <workspace>/_research/doc-fixtures/。
本脚本只用于**造测试数据**，不属于插件运行时依赖：
插件本身（lib/zip.js、lib/ooxml.js）只用 Node 内置模块。
"""

import datetime
import os
import sys
import zipfile
from pathlib import Path

HERE = Path(__file__).resolve().parent
DEFAULT_OUT = HERE.parents[1] / "_research" / "doc-fixtures"
OUT = Path(sys.argv[1]).resolve() if len(sys.argv) > 1 else DEFAULT_OUT

W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
CT_NS = "http://schemas.openxmlformats.org/package/2006/content-types"
REL_NS = "http://schemas.openxmlformats.org/package/2006/relationships"

COMMENTS_XML = f"""<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<w:comments xmlns:w="{W_NS}">
  <w:comment w:id="1" w:author="审阅者" w:initials="S" w:date="2024-01-01T00:00:00Z">
    <w:p><w:r><w:t>第一条批注：这里的数据需要复核。</w:t></w:r></w:p>
  </w:comment>
  <w:comment w:id="2" w:author="审阅者" w:initials="S" w:date="2024-01-02T00:00:00Z">
    <w:p><w:r><w:t>第二条批注：结论要更谨慎。</w:t></w:r></w:p>
  </w:comment>
</w:comments>
"""

FOOTNOTES_XML = f"""<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<w:footnotes xmlns:w="{W_NS}">
  <w:footnote w:type="separator" w:id="-1"><w:p><w:r><w:separator/></w:r></w:p></w:footnote>
  <w:footnote w:id="1"><w:p><w:r><w:t>脚注一：数据来源见附录。</w:t></w:r></w:p></w:footnote>
  <w:footnote w:id="2"><w:p><w:r><w:t>脚注二：样本量偏小。</w:t></w:r></w:p></w:footnote>
</w:footnotes>
"""

# 修订：删除的文字不能出现在抽取结果里，插入的文字必须出现。
REVISIONS_XML = f"""<w:p><w:r><w:t>修订演示：</w:t></w:r><w:del w:id="9" w:author="审阅者" w:date="2024-01-01T00:00:00Z"><w:r><w:delText>被删除的文字</w:delText></w:r></w:del><w:ins w:id="10" w:author="审阅者" w:date="2024-01-01T00:00:00Z"><w:r><w:t>插入后的文字</w:t></w:r></w:ins></w:p>"""


def rewrite_zip(path, edits):
    """按 {部件名: 新文本} 重写 zip：已有部件替换，新增部件追加，其余原样保留。"""
    with zipfile.ZipFile(path) as zin:
        items = [(i.filename, zin.read(i.filename)) for i in zin.infolist()]
    seen = set()
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as zout:
        for name, data in items:
            if name in edits:
                data = edits[name].encode("utf-8")
                seen.add(name)
            zout.writestr(name, data)
        for name, text in edits.items():
            if name not in seen:
                zout.writestr(name, text.encode("utf-8"))


def build_docx(path, png_path):
    from docx import Document
    from docx.enum.text import WD_BREAK
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    doc = Document()
    doc.add_heading("文档抽取测试 Document Extraction Test", level=1)
    doc.add_paragraph("中文与 English 混排：模型上下文窗口 test 123。")
    doc.add_paragraph("")  # 空段落
    p = doc.add_paragraph()
    p.add_run("段落A")
    p.add_run("同一段落内的第二段文字。")

    p = doc.add_paragraph()
    r = p.add_run("制表符→")
    r.add_tab()
    r.add_text("列二")
    r.add_tab()
    r.add_text("列三")

    p = doc.add_paragraph()
    r = p.add_run("软换行前")
    r.add_break(WD_BREAK.LINE)
    r.add_text("软换行后")

    p = doc.add_paragraph()
    r = p.add_run("特殊符号：不断行连字符")
    sym = OxmlElement("w:noBreakHyphen")
    r._r.append(sym)  # w:noBreakHyphen → "-"
    r2 = p.add_run("，项目符号")
    symbol = OxmlElement("w:sym")
    symbol.set(qn("w:font"), "Wingdings")
    symbol.set(qn("w:char"), "F0B7")  # → "·"
    r2._r.append(symbol)

    long_text = "超长段落：" + ("这是用来测试窗口读取与截断的重复内容 ABCdef123。" * 120)
    doc.add_paragraph(long_text)

    table = doc.add_table(rows=3, cols=3)
    data = [["表头一", "表头二", "表头三"], ["甲", "乙", "丙"], ["a1", "b2", "c3"]]
    for i, row in enumerate(data):
        for j, text in enumerate(row):
            table.cell(i, j).text = text
    table.cell(1, 0).add_paragraph("单元格第二段")

    doc.add_page_break()
    doc.add_paragraph("分页后的段落。")

    section = doc.sections[0]
    if section.header.paragraphs:
        section.header.paragraphs[0].text = "页眉：内部资料 Header"
    else:
        section.header.add_paragraph("页眉：内部资料 Header")
    if section.footer.paragraphs:
        section.footer.paragraphs[0].text = "页脚：第 1 页 Footer"
    else:
        section.footer.add_paragraph("页脚：第 1 页 Footer")

    doc.add_picture(str(png_path))
    doc.save(path)

    # 后处理：注入批注部件、脚注部件、修订标记（python-docx 没有这些 API）
    with zipfile.ZipFile(path) as zin:
        document = zin.read("word/document.xml").decode("utf-8")
        content_types = zin.read("[Content_Types].xml").decode("utf-8")
        doc_rels = zin.read("word/_rels/document.xml.rels").decode("utf-8")

    overrides = (
        '<Override PartName="/word/comments.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.comments+xml"/>'
        '<Override PartName="/word/footnotes.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.footnotes+xml"/>'
    )
    content_types = content_types.replace("</Types>", overrides + "</Types>")

    rels_add = (
        f'<Relationship Id="rIdComments" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/comments" Target="comments.xml"/>'
        f'<Relationship Id="rIdFootnotes" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/footnotes" Target="footnotes.xml"/>'
    )
    doc_rels = doc_rels.replace("</Relationships>", rels_add + "</Relationships>")

    document = document.replace("</w:body>", REVISIONS_XML + "</w:body>")

    rewrite_zip(
        path,
        {
            "word/document.xml": document,
            "[Content_Types].xml": content_types,
            "word/_rels/document.xml.rels": doc_rels,
            "word/comments.xml": COMMENTS_XML,
            "word/footnotes.xml": FOOTNOTES_XML,
        },
    )


def build_pptx(path):
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])  # 标题 + 内容
    slide.shapes.title.text = "第一页：抽取测试"
    body = slide.placeholders[1].text_frame
    body.text = "要点一：中文与 English 混排"
    para = body.add_paragraph()
    para.text = "要点二：第二个段落"
    para = body.add_paragraph()
    para.text = "要点三：带\t制表符"
    slide.notes_slide.notes_text_frame.text = "备注：第一页的讲稿。第二句。"

    slide2 = prs.slides.add_slide(prs.slide_layouts[5])  # 仅标题
    slide2.shapes.title.text = "第二页 Slide Two"
    table = slide2.shapes.add_table(2, 2, Inches(1), Inches(2), Inches(6), Inches(1.5)).table
    table.cell(0, 0).text = "表头A"
    table.cell(0, 1).text = "表头B"
    table.cell(1, 0).text = "值1"
    table.cell(1, 1).text = "值2"
    slide2.notes_slide.notes_text_frame.text = "备注：第二页讲稿。"

    # 第三页故意不加备注，用于验证备注计数
    slide3 = prs.slides.add_slide(prs.slide_layouts[5])
    slide3.shapes.title.text = "第三页：没有备注"

    prs.save(str(path))


def build_xlsx(path):
    """XlsxWriter：共享字符串表 + 公式缓存 + 稀疏列 + 日期序列号。"""
    import xlsxwriter

    wb = xlsxwriter.Workbook(str(path))
    ws = wb.add_worksheet("数据")
    date_fmt = wb.add_format({"num_format": "yyyy-mm-dd"})

    ws.write("A1", "名称")
    ws.write("B1", "数值")
    ws.write("C1", "公式")
    ws.write("E1", "稀疏列E")
    ws.write("A2", "中文条目一")
    ws.write_number("B2", 42)
    ws.write_formula("C2", "=B2*2")
    ws.write("A3", "English row")
    ws.write_number("B3", 3.14)
    ws.write_formula("C3", "=SUM(B2:B3)")
    ws.write_datetime("B4", datetime.datetime(2024, 3, 15), date_fmt)  # 存成序列号
    ws.write("A5", "布尔")
    ws.write_boolean("B5", True)
    ws.write("A6", "错误值")
    ws.write_formula("B6", "=1/0")  # 缓存值 -> t="e"
    ws.write_formula("C6", '=CONCATENATE("a","b")', None, "ab")  # 缓存串 -> t="str"
    ws.write("A8", "上面第 7 行是空行")

    ws2 = wb.add_worksheet("汇总")
    ws2.write("A1", "汇总表")
    ws2.write("A2", 100)
    ws2.write("B2", "引用数据表")
    wb.close()

    # XlsxWriter 对 =1/0 只写缓存的 0；Excel 重算后写的是错误值，这里改成 Excel 的写法，
    # 用来覆盖 t="e"（错误值）分支。
    with zipfile.ZipFile(path) as zin:
        sheet = zin.read("xl/worksheets/sheet1.xml").decode("utf-8")
    old = '<c r="B6"><f>1/0</f><v>0</v></c>'
    assert old in sheet, "XlsxWriter 的 B6 单元格写法变了，请更新 make-fixtures.py"
    rewrite_zip(path, {"xl/worksheets/sheet1.xml": sheet.replace(old, '<c r="B6" t="e"><v>#DIV/0!</v></c>')})


def build_inline_xlsx(path):
    """openpyxl：字符串写成 inlineStr（覆盖 t="inlineStr" 分支）。"""
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    ws = wb.active
    ws.title = "内联串"
    ws["A1"] = "内联字符串 inline"
    ws["B1"] = 123
    ws["C1"] = "中文 inlineStr 第三列"
    ws["A3"] = "稀疏：这一行只有 A 和 D"
    ws["D3"] = 7
    # 第 5 行只有格式没有值 —— 会写出 <row r="5"> 但没有内容，用来测「跳过空行」
    ws.cell(row=5, column=2).font = Font(bold=True)
    ws2 = wb.create_sheet("第二表")
    ws2["A1"] = 3.5
    ws2["A2"] = "第二表的中文"
    ws2.sheet_state = "hidden"  # 测隐藏工作表的提示
    wb.save(str(path))


def build_many_sheets(path, count=25):
    from openpyxl import Workbook

    wb = Workbook()
    wb.active.title = "表1"
    for i in range(2, count + 1):
        wb.create_sheet(f"表{i}")
    wb["表1"]["A1"] = "第一张表"
    wb[f"表{count}"]["A1"] = "最后一张表"
    wb.save(str(path))


def build_many_rows(path, count=50):
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "长表"
    for i in range(1, count + 1):
        ws.cell(row=i, column=1, value=f"第{i}行")
        ws.cell(row=i, column=2, value=i)
    wb.save(str(path))


def build_oracle_texts(out_dir):
    """用 python-docx / python-pptx / openpyxl 抽出同样的文本，作为人工比对的 oracle。"""
    from docx import Document
    from docx.table import Table
    from docx.text.paragraph import Paragraph
    from pptx import Presentation
    from openpyxl import load_workbook

    docx_path = out_dir / "sample.docx"
    doc = Document(str(docx_path))
    body = doc.element.body
    lines = []
    for child in body.iterchildren():
        tag = child.tag.split("}")[-1]
        if tag == "p":
            lines.append(Paragraph(child, doc).text)
        elif tag == "tbl":
            table = Table(child, doc)
            for row in table.rows:
                lines.append(" | ".join(cell.text.replace("\n", " ") for cell in row.cells))
    lines.append("===== 页眉/页脚 =====")
    for section in doc.sections:
        for part, label in ((section.header, "header"), (section.footer, "footer")):
            for p in part.paragraphs:
                lines.append(f"[{label}] {p.text}")
    (out_dir / "oracle-docx.txt").write_text("\n".join(lines), encoding="utf-8")

    pptx_path = out_dir / "sample.pptx"
    prs = Presentation(str(pptx_path))
    plines = []
    for idx, slide in enumerate(prs.slides, start=1):
        plines.append(f"===== slide {idx} =====")
        for shape in slide.shapes:
            if shape.has_text_frame:
                for p in shape.text_frame.paragraphs:
                    if p.text.strip():
                        plines.append(p.text)
            if getattr(shape, "has_table", False) and shape.has_table:
                for row in shape.table.rows:
                    for cell in row.cells:
                        for p in cell.text_frame.paragraphs:
                            if p.text.strip():
                                plines.append(p.text)
        if slide.has_notes_slide:
            plines.append("[备注]")
            for p in slide.notes_slide.notes_text_frame.paragraphs:
                if p.text.strip():
                    plines.append(p.text)
    (out_dir / "oracle-pptx.txt").write_text("\n".join(plines), encoding="utf-8")

    for name in ("sample.xlsx", "sample-inline.xlsx"):
        for data_only in (False, True):
            wb = load_workbook(str(out_dir / name), data_only=data_only)
            xlines = []
            for ws in wb.worksheets:
                xlines.append(f'===== sheet "{ws.title}" =====')
                for row in ws.iter_rows():
                    cells = []
                    for cell in row:
                        value = cell.value
                        cells.append("" if value is None else str(value))
                    while cells and cells[-1] == "":
                        cells.pop()
                    if any(c != "" for c in cells):
                        xlines.append("\t".join(cells))
            suffix = "-values" if data_only else ""
            (out_dir / f"oracle-{name.replace('.xlsx', '')}{suffix}.txt").write_text(
                "\n".join(xlines), encoding="utf-8"
            )


def main():
    OUT.mkdir(parents=True, exist_ok=True)
    from PIL import Image

    png = OUT / "_pixel.png"
    Image.new("RGB", (48, 48), (200, 30, 30)).save(str(png))

    build_docx(OUT / "sample.docx", png)
    build_pptx(OUT / "sample.pptx")
    build_xlsx(OUT / "sample.xlsx")
    build_inline_xlsx(OUT / "sample-inline.xlsx")
    build_many_sheets(OUT / "many-sheets.xlsx")
    build_many_rows(OUT / "many-rows.xlsx")
    build_oracle_texts(OUT)

    for item in sorted(OUT.iterdir()):
        print(f"{item.name}\t{item.stat().st_size}")


if __name__ == "__main__":
    main()
